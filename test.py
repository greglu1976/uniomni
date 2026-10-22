from core.OrderHandler import OrderHandler
from core.MainConfigHandler import MainConfigHandler
import os, re
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


order_data = OrderHandler()
config = MainConfigHandler.from_json_file("meta.json")


def build_flat_structure(nodes, path=None):
    if path is None:
        path = []
    result = []
    for node in nodes or []:
        t = node.get("Type")
        name = node.get("Name")

        if t == "Group":
            result.extend(build_flat_structure(node.get("Nodes", []), path + [name]))

        elif t == "Parameter" and name:
            info = config.get_param_info(name) or {}
            result.append({
                "path": path,
                "node_name": name,
                "description": info.get("description", "") or "",
                "fullDescription": info.get("fullDescription", "") or "",
                "appliedDescription": info.get("appliedDescription", "") or "",
            })

        elif t == "VirtualKey":
            # У VirtualKey нет Nodes — берём параметры из ButtonParameter/LedParameter
            led_param = node.get("LedParameter") or node.get("ButtonParameter") or ""
            info = config.get_param_info(led_param) if led_param else None

            # Если в meta.json нет — заполняем тем, что есть
            if info:
                desc = info.get("description", "") or ""
                full_desc = info.get("fullDescription", "") or ""
                applied = info.get("appliedDescription", "") or ""
            else:
                desc = full_desc = applied = ""

            result.append({
                "path": path,
                "node_name": name or led_param or "",
                "description": desc,
                "fullDescription": full_desc,
                "appliedDescription": applied or led_param,  # fallback на имя LedParameter
            })

    return result


def sanitize(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip()


def export_trees_to_excel(tree_names, output_name, out_dir="output"):
    os.makedirs(out_dir, exist_ok=True)

    all_items = []
    for tree_name in tree_names:
        nodes = order_data.get_signals_by_name(tree_name)
        if not nodes:
            print(f"⚠  [{tree_name}] не найдено")
            continue
        items = build_flat_structure(nodes)
        all_items.extend(items)
        print(f"→  [{tree_name}] {len(items)} параметров")

    if not all_items:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sanitize(output_name)[:31]

    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    level_fills = {
        1: PatternFill("solid", fgColor="FCE4D6"),
        2: PatternFill("solid", fgColor="FFF2CC"),
        3: PatternFill("solid", fgColor="E2EFDA"),
        4: PatternFill("solid", fgColor="DDEBF7"),
    }

    ws["A1"] = output_name
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    for i, title in enumerate(["Наименование", "Обозначение ФС", "Полное описание", "Описание"], 1):
        c = ws.cell(row=3, column=i, value=title)
        c.font = bold
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    ws.freeze_panes = "A4"

    row = 4
    group_rows = {}
    for it in all_items:
        path = it["path"]
        for depth, gname in enumerate(path, start=1):
            key = tuple(path[:depth])
            if key in group_rows:
                continue
            c = ws.cell(row=row, column=1, value=gname)
            c.font = bold
            c.fill = level_fills.get(depth, level_fills[4])
            c.border = border
            for col in range(2, 5):
                cc = ws.cell(row=row, column=col)
                cc.fill = level_fills.get(depth, level_fills[4])
                cc.border = border
            ws.row_dimensions[row].outlineLevel = depth - 1
            group_rows[key] = row
            row += 1

        for col, val in enumerate([
            it["node_name"],
            it["appliedDescription"] or it["description"],
            it["fullDescription"],
            it["description"],
        ], 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = border
            c.alignment = wrap
        ws.row_dimensions[row].outlineLevel = len(path) if path else 1
        row += 1

    for i, w in enumerate([40, 30, 40, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_properties.outlinePr.summaryBelow = False

    path_out = os.path.join(out_dir, f"{sanitize(output_name)}.xlsx")
    wb.save(path_out)
    print(f"   ✓ {path_out}")


# =========================================================
# ЧТО КУДА
# =========================================================
trees_to_export = {
    "Настройка регистрации":   ["DigitalSignalsTree"],
    "Уставки РЗиА":            ["SettingsTree"],
    "Информация":              ["InformationTree"],
    "Конфигурация":            ["ConfigurationTree"],
    #"Матрица входов":          ["ParametersToHardwareDigitalInputsTree"],
    #"Матрица выходных реле":   ["ParametersToHardwareDigitalOutputsTree"],
        "Матрица входов и выходных реле":   [
        "ParametersToHardwareDigitalInputsTree",
        "ParametersToHardwareDigitalOutputsTree",
    ],
    "Настройка светодиодов и ФК":   [
        "VirtualKeyParametersTree",
        "ParametersToHmiButtonsTree",
        "ParametersToHmiSignalizationTree",
    ],
}

for output_name, tree_names in trees_to_export.items():
    print(f"\n=== {output_name} ===")
    try:
        export_trees_to_excel(tree_names, output_name, out_dir="output")
    except Exception as e:
        print(f"✗  [{output_name}] ошибка: {e}")

print("\nГотово.")