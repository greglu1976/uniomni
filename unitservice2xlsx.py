# =========================================================
# unitservice2xlsx.py
# Автопоиск в папке:
#   <NNNN>.json             → meta.json
#   <NNNN>_GROUPING.json    → grouping.json
# Экспорт деревьев в Excel.
# =========================================================

import os
import sys
import re

from core.OrderHandler import OrderHandler
from core.MainConfigHandler import MainConfigHandler

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def get_base_dir():
    # PyInstaller .exe → папка рядом с .exe
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    # обычный .py → папка со скриптом
    return os.path.dirname(os.path.abspath(__file__))


BASE_DIR = get_base_dir()


# =========================================================
# 0. АВТОПОИСК ФАЙЛОВ
# =========================================================

def find_files(base_dir: str):
    """
    Ищет в папке пару:
      <NNNN>.json             → meta.json
      <NNNN>_GROUPING.json    → grouping.json
    Возвращает dict: {"meta": path, "grouping": path}.
    """
    candidates = [f for f in os.listdir(base_dir) if f.lower().endswith(".json")]

    meta_re = re.compile(r"^(\d+)\.json$", re.IGNORECASE)
    grp_re  = re.compile(r"^(\d+)[_\-]?grouping\.json$", re.IGNORECASE)

    meta_matches, grp_matches = {}, {}
    for f in candidates:
        m = meta_re.match(f)
        if m:
            meta_matches[m.group(1)] = os.path.join(base_dir, f)
            continue
        m = grp_re.match(f)
        if m:
            grp_matches[m.group(1)] = os.path.join(base_dir, f)
            continue

    common = sorted(set(meta_matches) & set(grp_matches))
    if not common:
        raise FileNotFoundError(
            "Не найдена пара '<NNNN>.json' + '<NNNN>_GROUPING.json' "
            f"в папке: {base_dir}\n"
            f"  meta:     {sorted(meta_matches)}\n"
            f"  grouping: {sorted(grp_matches)}"
        )

    if len(common) > 1:
        print(f"⚠  Найдено несколько пар {common}, использую первую: {common[0]}")

    key = common[0]
    return {
        "meta":     meta_matches[key],
        "grouping": grp_matches[key],
    }


FILES = find_files(BASE_DIR)

print(f"BASE_DIR       = {BASE_DIR}")
print(f"meta.json      → {os.path.basename(FILES['meta'])}")
print(f"grouping.json  → {os.path.basename(FILES['grouping'])}")
print()


# =========================================================
# 1. ЗАГРУЗКА
# =========================================================

order_data = OrderHandler(
    grouping_path=FILES["grouping"],
    meta_path=FILES["meta"],
)
config = order_data.config_handler   # тот же MainConfigHandler


# =========================================================
# 2. СБОР ПЛОСКОЙ СТРУКТУРЫ
# =========================================================

def build_flat_structure(nodes, path=None):
    """
    Рекурсивно обходит дерево, возвращает плоский список параметров.
    path — список имён групп от корня до текущей.
    """
    if path is None:
        path = []
    result = []
    for node in nodes or []:
        t = node.get("Type")
        name = node.get("Name")

        if t == "Group":
            result.extend(build_flat_structure(node.get("Nodes", []), path + [name]))

        elif t == "Parameter" and name:
            info = config.get_param_info(name) or {}
            result.append({
                "path": path,
                "node_name": name,
                "description": info.get("description", "") or "",
                "fullDescription": info.get("fullDescription", "") or "",
                "appliedDescription": info.get("appliedDescription", "") or "",
            })

        elif t == "VirtualKey":
            # У VirtualKey нет Nodes — берём параметры из Led/Button
            led_param = node.get("LedParameter") or node.get("ButtonParameter") or ""
            info = config.get_param_info(led_param) if led_param else None
            if info:
                desc = info.get("description", "") or ""
                full_desc = info.get("fullDescription", "") or ""
                applied = info.get("appliedDescription", "") or ""
            else:
                desc = full_desc = applied = ""
            result.append({
                "path": path,
                "node_name": name or led_param or "",
                "description": desc,
                "fullDescription": full_desc,
                "appliedDescription": applied or name or led_param,
            })

    return result


def sanitize(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name).strip()


# =========================================================
# 3. ЭКСПОРТ В EXCEL
# =========================================================

def export_trees_to_excel(tree_names, output_name, out_dir="output"):
    os.makedirs(out_dir, exist_ok=True)

    all_items = []
    for tree_name in tree_names:
        nodes = order_data.get_signals_by_name(tree_name)
        if not nodes:
            print(f"⚠  [{tree_name}] не найдено")
            continue
        items = build_flat_structure(nodes)
        all_items.extend(items)
        print(f"→  [{tree_name}] {len(items)} параметров")

    if not all_items:
        print(f"⚠  [{output_name}] нет данных — файл не создан")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sanitize(output_name)[:31]

    bold = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")
    level_fills = {
        1: PatternFill("solid", fgColor="FCE4D6"),
        2: PatternFill("solid", fgColor="FFF2CC"),
        3: PatternFill("solid", fgColor="E2EFDA"),
        4: PatternFill("solid", fgColor="DDEBF7"),
    }

    # Заголовок листа
    ws["A1"] = output_name
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    # Шапка
    for i, title in enumerate(
        ["Наименование", "Обозначение ФС", "Полное описание", "Описание"], 1
    ):
        c = ws.cell(row=3, column=i, value=title)
        c.font = bold
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    ws.freeze_panes = "A4"

    row = 4
    group_rows = {}   # key = tuple(path[:depth]) → номер строки-заголовка

    for it in all_items:
        path = it["path"]

        # Заголовки групп (многоуровневые, с outlineLevel)
        for depth, gname in enumerate(path, start=1):
            key = tuple(path[:depth])
            if key in group_rows:
                continue

            c = ws.cell(row=row, column=1, value=gname)
            c.font = bold
            c.fill = level_fills.get(depth, level_fills[4])
            c.border = border
            for col in range(2, 5):
                cc = ws.cell(row=row, column=col)
                cc.fill = level_fills.get(depth, level_fills[4])
                cc.border = border

            ws.row_dimensions[row].outlineLevel = depth - 1
            group_rows[key] = row
            row += 1

        # Строка параметра
        for col, val in enumerate([
            it["node_name"],
            it["appliedDescription"] or it["description"],
            it["fullDescription"],
            it["description"],
        ], 1):
            c = ws.cell(row=row, column=col, value=val)
            c.border = border
            c.alignment = wrap

        ws.row_dimensions[row].outlineLevel = len(path) if path else 1
        row += 1

    # Ширины колонок
    for i, w in enumerate([40, 30, 40, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Кнопки +/- сверху группы
    ws.sheet_properties.outlinePr.summaryBelow = False

    path_out = os.path.join(out_dir, f"{sanitize(output_name)}.xlsx")
    wb.save(path_out)
    print(f"   ✓ {path_out}")


# =========================================================
# 4. ЧТО КУДА (порядок важен!)
# =========================================================

trees_to_export = {
    "Настройка регистрации":         ["DigitalSignalsTree"],
    "Уставки РЗиА":                  ["SettingsTree"],
    "Информация":                    ["InformationTree"],
    "Конфигурация":                  ["ConfigurationTree"],
    "Матрица входов и выходных реле": [
        "ParametersToHardwareDigitalInputsTree",
        "ParametersToHardwareDigitalOutputsTree",
    ],
    "Настройка светодиодов и ФК":    [
        "VirtualKeyParametersTree",
        "ParametersToHmiButtonsTree",
        "ParametersToHmiSignalizationTree",
    ],
}


# =========================================================
# 5. ОСНОВНОЙ ЦИКЛ
# =========================================================

def main():
    for output_name, tree_names in trees_to_export.items():
        print(f"\n=== {output_name} ===")
        try:
            export_trees_to_excel(tree_names, output_name, out_dir="output")
        except Exception as e:
            print(f"✗  [{output_name}] ошибка: {e}")
    print("\nГотово.")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n✗ Критическая ошибка: {e}")
    try:
        input("\nНажмите Enter, чтобы закрыть...")
    except EOFError:
        pass