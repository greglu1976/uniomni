# SettingBlanc2.py
# Формирование бланка уставок в формате docx.

import ast
import json
import os
import re
import tempfile

from docx import Document
from docxtpl import DocxTemplate, RichText

from core.ExtensionHandler import ExtensionHandler
from core.FsuInformation import FsuInformation
from core.MainConfigHandler import MainConfigHandler
from core.OrderHandler import OrderHandler
from core.UnitConverter import (DEFAULT_TIME_UNIT, DEFAULT_UNIT_TYPE,
                                SUFFIX_MARKER_WORD, with_suffix)
from logger.logger import Logger
from utils.docx_handler import (
    add_new_section,
    add_new_section_landscape,
    add_paragraph,
)
from utils.paths import find_file, project_path
from utils.tables import (
    ALIGN_CENTER,
    add_table_binaries_core4,
    add_table_final,
    add_table_fks_core4,
    add_table_leds_new_core4,
    add_table_mtrx_ins_core4,
    add_table_mtrx_outs_core4,
    add_table_reg_core4,
    add_table_settings_core4,
    row_cells,
    set_cell_multiline_text,
    set_cell_text,
    set_paragraph_alignment,
    set_runs_font_size,
)

STYLE_HEADING_1 = 'ДОК Заголовок 1'
STYLE_HEADING_2 = 'ДОК Заголовок 2'
STYLE_HEADING_3 = 'ДОК Заголовок 3'
STYLE_TABLE_TITLE = 'ДОК Таблица Название'
STYLE_TAGS = 'TAGS'

# Регулярное выражение для последней группы в круглых скобках
_PARENTHESIS_RE = re.compile(r'\(([^()]*)\)')


class SettingBlanc:
    """Бланк уставок: docx-документ с таблицами уставок, входов/выходов и т.д."""

    def __init__(self, device_data, unit_type=DEFAULT_UNIT_TYPE,
                 time_unit=DEFAULT_TIME_UNIT, convert_step=True, use_fsu=True):
        packet_path = device_data.get('path_to_support_packets')
        Logger.info(f"Путь к пакетам поддержки: {packet_path}")

        self.device_data = device_data
        self.unit_type = unit_type or DEFAULT_UNIT_TYPE
        self.time_unit = time_unit or DEFAULT_TIME_UNIT
        self.convert_step = bool(convert_step)
        # Источник уставок: fsu-information.json (по умолчанию) или GROUPING/meta
        self.use_fsu = bool(use_fsu)
        self.fsu_information = None
        self.code = self.device_data.get("setting_blanc_code", "")
        self.versions = self.device_data.get("versions") or []
        self.base_structure = None

        if not packet_path:
            raise ValueError("В конфигурации устройства не задан path_to_support_packets")

        self.extension_handler = ExtensionHandler(packet_path)

        meta_path = find_file(packet_path, "meta.json")
        if meta_path is None:
            raise FileNotFoundError(f"В пакете поддержки не найден meta.json: {packet_path}")
        self.config_handler = MainConfigHandler.from_json_file(str(meta_path))

        self.order_handler = OrderHandler(self.config_handler, self.extension_handler,
                                          packet_path, unit_type=self.unit_type,
                                          time_unit=self.time_unit,
                                          convert_step=self.convert_step)

        if self.use_fsu:
            self.fsu_information = FsuInformation.try_load(packet_path)
            if self.fsu_information is None:
                self.use_fsu = False

        self.di_list = []
        self.maps = self.order_handler.get_mapping()

        abbr_path = project_path("abbr.json")
        with open(abbr_path, 'r', encoding='utf-8') as file:
            self.abbr_dict = json.load(file)

    # ------------------------------------------------------------------ утилиты

    def _make_richtext(self, text):
        if not text:
            return RichText()
        text = text.replace('\\n', '\n')
        rich_text = RichText()
        for i, part in enumerate(text.split('\n')):
            if i > 0:
                rich_text.add('\n')
            rich_text.add(part)
        return rich_text

    @staticmethod
    def _split_parenthesis(raw_text):
        """Отделяет последнюю группу в круглых скобках от наименования.

        Пример: 'Начальный ток (Iср_загр)' → ('Начальный ток', 'Iср_загр').
        """
        raw_text = "" if raw_text is None else str(raw_text)
        matches = list(_PARENTHESIS_RE.finditer(raw_text))
        if not matches:
            return raw_text.strip(), ""

        last_match = matches[-1]
        name = (raw_text[:last_match.start()] + raw_text[last_match.end():]).strip()
        return name, last_match.group(1).strip()

    def _display_name(self, raw_col1):
        """Наименование уставки для таблицы.

        В meta.json краткое обозначение идёт в скобках в конце; оно терялось.
        Возвращаем его, только если оно не повторяет уже имеющийся текст.
        """
        name, short_name = self._split_parenthesis(raw_col1)
        if short_name and short_name not in name:
            return f"{name} ({short_name})"
        return name

    @staticmethod
    def _parse_note_dict(note_str):
        """Преобразует 'note_{"0": "a", "1": "b"}' в 'a /\\nb'."""
        if not isinstance(note_str, str) or not note_str.startswith('note_'):
            return note_str

        try:
            note_dict = ast.literal_eval(note_str[5:])
            if not isinstance(note_dict, dict):
                return note_str
            # Ключи могут быть как числовыми, так и текстовыми — сортируем
            # их без смешивания типов (иначе TypeError).
            sorted_items = sorted(
                note_dict.items(),
                key=lambda item: (0, int(item[0]))
                if str(item[0]).isdigit()
                else (1, str(item[0])),
            )
            return ' /\n'.join(str(value) for _, value in sorted_items)
        except (ValueError, SyntaxError, TypeError) as error:
            Logger.warning(f"Ошибка разбора значений уставки: {error}")
            return note_str

    # --------------------------------------------------------------- уставки

    def _create_section_settings_core4(self, doc):
        if not self.base_structure:
            Logger.warning("Нет данных base_structure для генерации уставок")
            return

        add_new_section_landscape(doc)
        add_paragraph(doc, 'УСТАВКИ РЗиА', STYLE_HEADING_1)

        for func_block in self.base_structure:
            func_type = func_block.get('type')
            func_name = func_block.get('func_name', 'Без имени')
            fixed_func_name = func_name.split("_")[0]

            if "fsu-information" in fixed_func_name:
                fixed_func_name = "Общие уставки"

            title = self.abbr_dict.get(fixed_func_name, fixed_func_name)

            if func_type == 'simple':
                add_paragraph(doc, title, STYLE_HEADING_2)
                add_paragraph(doc, fixed_func_name, STYLE_TABLE_TITLE)

                if func_block.get('rows'):
                    table = add_table_settings_core4(doc)
                    self._fill_table_settings(table, func_block['rows'])

            elif func_type == 'complex':
                add_paragraph(doc, title, STYLE_HEADING_2)

                # Собственные уставки блока (источник fsu-information):
                # в GROUPING такие уставки не выводились вовсе
                if func_block.get('rows'):
                    add_paragraph(doc, title, STYLE_TABLE_TITLE)
                    table = add_table_settings_core4(doc)
                    self._fill_table_settings(table, func_block['rows'])
                    add_paragraph(doc, '', STYLE_TAGS)

                for sub_func in func_block.get('sub_functions', []):
                    subtitle = sub_func.get('subtitle', '')
                    if subtitle:
                        add_paragraph(doc, subtitle, STYLE_TABLE_TITLE)

                    if sub_func.get('rows'):
                        table = add_table_settings_core4(doc)
                        self._fill_table_settings(table, sub_func['rows'])
                        add_paragraph(doc, '', STYLE_TAGS)

    def _fill_table_settings(self, table, rows_data):
        """Заполняет таблицу уставок строками (по 10 столбцов)."""
        for i, row_data in enumerate(rows_data, start=1):
            row = table.add_row()
            cells = row_cells(row, table)

            parameter_name = row_data.get('col0')
            enum500 = self.extension_handler.find_enum_by_parameter_name(parameter_name)

            display_name = self._display_name(row_data.get('col1', ''))

            set_cell_text(cells[0], i)
            set_paragraph_alignment(cells[0].paragraphs[0], ALIGN_CENTER)

            set_cell_multiline_text(cells[1], display_name)

            suffix = row_data.get('suffix')

            col3_value = row_data.get('col3', '')
            if isinstance(col3_value, str) and col3_value.startswith('note_'):
                col3_value = self._parse_note_dict(col3_value)
            if enum500:
                col3_value = " / ".join(str(item.get('VisibleValue', '')) for item in enum500)
            set_cell_text(cells[2], with_suffix(col3_value, suffix,
                                                SUFFIX_MARKER_WORD, wrap_range=True))
            set_paragraph_alignment(cells[2].paragraphs[0], ALIGN_CENTER)

            set_cell_text(cells[3], row_data.get('col4'))
            set_paragraph_alignment(cells[3].paragraphs[0], ALIGN_CENTER)

            set_cell_text(cells[4], with_suffix(row_data.get('col5'), suffix,
                                                SUFFIX_MARKER_WORD))
            set_paragraph_alignment(cells[4].paragraphs[0], ALIGN_CENTER)

            col6_value = row_data.get('col6', '')
            if enum500:
                lookup = {item.get('ParameterValue'): item.get('VisibleValue') for item in enum500}
                try:
                    col6_value = lookup.get(int(col6_value), col6_value)
                except (ValueError, TypeError):
                    pass
            set_cell_text(cells[5], with_suffix(col6_value, suffix,
                                                SUFFIX_MARKER_WORD))
            set_paragraph_alignment(cells[5].paragraphs[0], ALIGN_CENTER)

            for j in range(6, 10):
                set_cell_text(cells[j], '')

        set_runs_font_size(table, 10)
        return table

    def get_all_settings(self):
        """Собирает блоки уставок из выбранного источника (fsu или GROUPING)."""
        if self.use_fsu and self.fsu_information is not None:
            base_structure = self.order_handler.build_settings_blocks_from_fsu(
                self.fsu_information)
            # Для обновления таблиц в РЭ используем те же имена блоков
            self.maps = {block['func_name']: block['func_name'] for block in base_structure}
            Logger.info(f"Источник уставок: fsu-information.json "
                        f"({len(base_structure)} блоков)")
        else:
            base_structure = []
            for fb in list(self.maps.keys()):
                fb_map = self.maps.get(fb)
                if not fb_map:
                    continue
                json_data = self.order_handler.get_data_by_fb_name(fb_map)
                base_structure.extend(
                    self.order_handler.parse_rza_structure(json_data, all_struct=0))
            Logger.info(f"Источник уставок: GROUPING.json/meta.json "
                        f"({len(base_structure)} блоков)")

        self.base_structure = base_structure
        Logger.info(f"Загружено {len(base_structure)} блоков уставок")
        return base_structure

    # ------------------------------------------------------------- документ

    def create_template(self):
        versions = self.device_data.get('versions') or []
        last_version = versions[-1] if versions else {}
        edition = last_version.get('edition', '')
        release_date = last_version.get('data', '')
        colontile = f"Редакция {edition} от {release_date}" if versions else ''

        if not versions:
            Logger.warning("В конфигурации устройства не заданы версии — "
                           "колонтитул и имя файла будут без редакции")

        template_path = project_path('origin.docx')
        if not template_path.is_file():
            raise FileNotFoundError(f"Не найден шаблон бланка уставок: {template_path}")

        config_version = self.config_handler.config_version or ""
        if "-ЮНИТ-" in config_version:
            first_part, second_part = config_version.split("-ЮНИТ-", 1)
            second_part = "ЮНИТ-" + second_part
        else:
            first_part, second_part = config_version, None

        context = {
            "title": self._make_richtext(self.device_data.get('full_description')),
            "code": self.device_data.get('setting_blanc_code', ''),
            "device_order_code": first_part,
            "hmi_order_code": second_part,
            "versions": versions,
            "device_name": self.device_data.get('name', ''),
            "colontile": colontile,
            "packet": self.config_handler.model_version,
            "version": edition,
            "date": release_date,
        }

        doc_tpl = DocxTemplate(str(template_path))
        doc_tpl.render(context)

        # Промежуточный файл создаём во временном каталоге и сразу удаляем
        temp_fd, temp_path = tempfile.mkstemp(suffix='.docx')
        os.close(temp_fd)
        try:
            doc_tpl.save(temp_path)
            doc = Document(temp_path)
        finally:
            try:
                os.unlink(temp_path)
            except OSError:
                pass

        self.get_all_settings()
        Logger.info("Создаем раздел Уставки РЗиА...")
        self._create_section_settings_core4(doc)
        Logger.info("Создаем раздел Матрица входов и выходных реле...")
        self._create_section_inouts_core4(doc)

        if second_part:
            Logger.info("ИЧМ присутствует. Создаем раздел Настройка светодиодов и ФК...")
            self._create_section_leds_core4(second_part, doc)

        Logger.info("Создаем раздел Конфигурация...")
        self._create_section_config_core4(doc)
        Logger.info("Создаем раздел Настройка регистрации...")
        self._create_section_disturb_core4(doc)

        add_new_section(doc)
        add_table_final(doc)

        name_for_save = f"БУ {self.device_data.get('name', '')} ({self.code})"
        if edition:
            name_for_save += f"_v.{edition}"
        output_path = project_path(f'{name_for_save}.docx')
        doc.save(str(output_path))
        Logger.info(f"Бланк уставок сохранен: '{output_path}'")
        return doc

    def get_blanc(self):
        return self.create_template()

    # ------------------------------------------------------- входы/выходы

    @staticmethod
    def _extract_description(item):
        if isinstance(item, dict):
            return (item.get('fullDescription')
                    or item.get('appliedDescription')
                    or item.get('description')
                    or item.get('name', ''))
        return str(item)

    @staticmethod
    def _slot_sort_key(slot_name):
        match = re.search(r'M(\d+)', slot_name or '')
        return int(match.group(1)) if match else 0

    def _collect_slot_signals(self, items_to_process, pattern):
        """Собирает подписи статус-сигналов слотов по регулярному выражению."""
        result = {}
        for slot_name, params_list in items_to_process:
            status_signals = [param for param in params_list if pattern.match(param)]
            if not status_signals:
                continue

            descriptions = []
            for signal in status_signals:
                info = self.config_handler.get_param_info(signal) or {}
                description = info.get("appliedDescription") or signal
                descriptions.append(str(description).replace(". Статус", "").strip())

            if descriptions:
                result[slot_name] = descriptions
        return result

    def _create_section_inouts_core4(self, doc):
        raw_sigs, raw_gen_sigs = self.order_handler.get_fsu_signals()
        sigs_list = [desc for desc in (self._extract_description(s) for s in raw_sigs) if desc]
        gen_signs = [desc for desc in (self._extract_description(s) for s in raw_gen_sigs) if desc]

        slots_data = self.order_handler.get_slots_data()
        items_to_process = []
        if isinstance(slots_data, list):
            for slot_dict in slots_data:
                items_to_process.extend(slot_dict.items())
        elif isinstance(slots_data, dict):
            items_to_process.extend(slots_data.items())

        pattern_inputs = re.compile(r'^M\d+_B\d{3}_B\d+_Status$')
        pattern_outputs = re.compile(r'^M\d+_K\d{3}_K\d+_Status$')

        final_dict_inputs = self._collect_slot_signals(items_to_process, pattern_inputs)
        final_dict_outputs = self._collect_slot_signals(items_to_process, pattern_outputs)

        if not final_dict_inputs and not final_dict_outputs:
            return

        add_new_section_landscape(doc)
        add_paragraph(doc, 'МАТРИЦА ВХОДОВ И ВЫХОДНЫХ РЕЛЕ', STYLE_HEADING_1)

        if final_dict_inputs:
            add_paragraph(doc, 'Дискретные входы', STYLE_HEADING_2)
            for slot_name in sorted(final_dict_inputs, key=self._slot_sort_key):
                add_paragraph(doc, slot_name, STYLE_TABLE_TITLE)
                add_table_mtrx_ins_core4(
                    doc=doc,
                    slot_name=slot_name,
                    inputs_list=final_dict_inputs[slot_name],
                    sigs=sigs_list,
                    di_sigs=gen_signs,
                )
                doc.add_paragraph()

        if final_dict_outputs:
            add_paragraph(doc, 'Выходные реле', STYLE_HEADING_2)
            outs_list = self.extract_all_signals_from_structure(
                self.order_handler.get_fsu_out_signals())
            for slot_name in sorted(final_dict_outputs, key=self._slot_sort_key):
                add_paragraph(doc, slot_name, STYLE_TABLE_TITLE)
                add_table_mtrx_outs_core4(
                    doc=doc,
                    outputs_list=final_dict_outputs[slot_name],
                    sigs_list=outs_list,
                )
                doc.add_paragraph()

    def extract_all_signals_from_structure(self, data_structure):
        all_signals = []
        for function_item in data_structure or []:
            for subfunc in function_item.get('subfunctions', []):
                for signal in subfunc.get('data', []):
                    if signal:
                        all_signals.append(signal)
        return all_signals

    # ------------------------------------------------ светодиоды и клавиши

    @staticmethod
    def _extension_modules(order_code):
        """Расшифровывает позиции 4-7 кода заказа ИЧМ (модули расширения)."""
        parts = (order_code or '').split('-')
        result = []
        for number in range(1, 5):
            position = 2 + number  # parts[3]..parts[6]
            code = parts[position] if position < len(parts) else ''
            if code == 'С':
                result.append(f"Модуль расширения {number} на 16 светодиодов")
            elif code == 'К':
                result.append(f"Модуль расширения {number} на 16 функциональных кнопок")
            else:
                result.append("Модуль отсутствует")
        return result

    def _create_section_leds_core4(self, order_code, doc):
        modules = self._extension_modules(order_code)

        add_new_section_landscape(doc)
        add_paragraph(doc, 'НАСТРОЙКА СВЕТОДИОДОВ И ФУНКЦИОНАЛЬНЫХ КЛАВИШ', STYLE_HEADING_1)
        add_paragraph(doc, 'Светодиоды', STYLE_HEADING_2)

        drop_list = self.order_handler.get_digital_signals_for_led()
        add_paragraph(doc, "ИЧМ", STYLE_TABLE_TITLE)
        add_table_leds_new_core4(doc, drop_list)

        for description in modules:
            if "светодиодов" in description:
                doc.add_paragraph()
                add_paragraph(doc, description, STYLE_TABLE_TITLE)
                add_table_leds_new_core4(doc, drop_list)

        add_paragraph(doc, 'Функциональные клавиши', STYLE_HEADING_2)

        buttons_list = self.order_handler.get_fsu_hmi_buttons()
        add_paragraph(doc, "ИЧМ", STYLE_TABLE_TITLE)
        add_table_fks_core4(doc, buttons_list)

        for description in modules:
            if "функциональных кнопок" in description:
                doc.add_paragraph()
                add_paragraph(doc, description, STYLE_TABLE_TITLE)
                # Раньше сюда по ошибке передавался drop_list (сигналы светодиодов)
                add_table_fks_core4(doc, buttons_list)

    # ------------------------------------------------------ конфигурация

    @staticmethod
    def enum_calc(enum, default):
        """'a / b / c' и значение по умолчанию, соответствующее коду."""
        result_str = " / ".join(str(item.get('VisibleValue', '')) for item in enum)
        default_visible = None
        for item in enum:
            if str(item.get('ParameterValue')) == str(default):
                default_visible = item.get('VisibleValue')
                break
        return result_str, default_visible

    @staticmethod
    def parse_note(note, default):
        """Парсит '0 - Вывод, 1 - Ввод' в ('Вывод / Ввод', <значение по умолчанию>)."""
        values = []
        default_value = None
        for part in str(note).split(','):
            if ' - ' not in part:
                continue
            key_str, value = part.split(' - ', 1)
            value = value.strip()
            values.append(value)
            if key_str.strip() == str(default):
                default_value = value
        return " / ".join(values), default_value

    def _create_section_config_core4(self, doc):
        add_new_section(doc)
        add_paragraph(doc, 'КОНФИГУРАЦИЯ', STYLE_HEADING_1)

        for datum in self.order_handler.get_data_for_configuration():
            add_paragraph(doc, datum["main_title"], STYLE_HEADING_2)

            for table in datum["tables"]:
                add_paragraph(doc, table["title"], STYLE_TABLE_TITLE)

                fixed_rows = []
                for row in table["rows"]:
                    fixed_rows.append(self._config_row(row["name"]))

                add_table_binaries_core4(doc, fixed_rows)
                doc.add_paragraph()

    def _config_row(self, row_name):
        """Готовит строку таблицы «Конфигурация» по имени параметра."""
        row_data = self.config_handler.get_param_info(row_name)
        if not row_data:
            Logger.warning(f"Параметр конфигурации отсутствует в meta.json: {row_name}")
            return (row_name, '', '-', '-', '-', '-')

        col1 = row_data.get("fullDescription", "")
        col2 = row_data.get("appliedDescription", "")
        note = str(row_data.get("note") or "")
        col3 = note if note.count('-') >= 2 else ""
        col4 = row_data.get("units") or '-'
        col5 = row_data.get("step") or '-'
        col6 = row_data.get("defaultValue", "")

        min_value = row_data.get("minValue")
        max_value = row_data.get("maxValue")

        if col3:
            col3, col6 = self.parse_note(col3, col6)

        if str(col6) == "false":
            col3, col6 = 'Вывод / Ввод', "Вывод"
        if str(col6) == "true":
            col3, col6 = 'Вывод / Ввод', "Ввод"

        enum_data = self.extension_handler.find_enum_by_parameter_name(row_name)
        if enum_data:
            col3, col6 = self.enum_calc(enum_data, col6)
            col4 = col5 = '-'
        elif col3:
            col4 = col5 = '-'
        elif min_value is None:
            col3 = '-'
        else:
            if str(min_value) == "0" and str(max_value) == "1":
                col3 = 'Вывод / Ввод'
                col6 = "Вывод" if str(row_data.get("defaultValue")) == "0" else "Ввод"
                col5 = '-'
            else:
                col3 = f'{min_value} ... {max_value}'

        return (col1, col2, col3, col4, col5, col6)

    # -------------------------------------------------------- регистрация

    def _create_section_disturb_core4(self, doc):
        add_new_section_landscape(doc)
        add_paragraph(doc, 'НАСТРОЙКА РЕГИСТРАЦИИ', STYLE_HEADING_1)

        for section in self.order_handler.get_data_for_registration():
            add_paragraph(doc, section["main_title"], STYLE_HEADING_2)

            for subsection in section.get("subsections", []):
                short_name = subsection["title"].split("_")[0]
                add_paragraph(doc, self.abbr_dict.get(short_name, short_name), STYLE_HEADING_3)

                for table_data in subsection["tables"]:
                    add_paragraph(doc, table_data["title"], STYLE_TABLE_TITLE)

                    data_rows = []
                    for param_name in table_data["parameters"]:
                        row_info = self.config_handler.get_param_info(param_name)
                        if not row_info:
                            continue
                        data_rows.append((
                            row_info.get("fullDescription", ""),
                            row_info.get("appliedDescription", ""),
                            row_info.get("type"),
                        ))

                    if data_rows:
                        add_table_reg_core4(doc, data_rows)
                        add_paragraph(doc, '', STYLE_TAGS)

    # ------------------------------------------------- уставки для latex

    def get_table_settings_latex(self, fb_key):
        """Данные таблицы уставок для latex (структура из Manual)."""
        if not self.base_structure:
            self.get_all_settings()

        block = None
        for func_block in self.base_structure:
            if fb_key == func_block.get('func_name'):
                block = func_block
                break

        if not block:
            Logger.warning(f"Блок {fb_key} не найден")
            return None

        func_type = block.get('type', 'simple')

        def prepare_rows(rows_data):
            result_rows = []
            for row_data in rows_data:
                raw_col1 = "" if row_data.get('col1') is None else str(row_data.get('col1'))
                parameter_name = row_data.get('col0')
                enum500 = self.extension_handler.find_enum_by_parameter_name(parameter_name)

                name, short_name = self._split_parenthesis(raw_col1)

                col3_value = row_data.get('col3', '')
                if enum500:
                    col3_value = " / ".join(
                        str(item.get('VisibleValue', '')) for item in enum500)

                col6_value = row_data.get('col6', '')
                if enum500:
                    lookup = {item.get('ParameterValue'): item.get('VisibleValue')
                              for item in enum500}
                    try:
                        col6_value = lookup.get(int(col6_value), col6_value)
                    except (ValueError, TypeError):
                        pass

                result_rows.append({
                    "col1": name,
                    "col2": short_name,
                    "display_name": self._display_name(raw_col1),
                    "col3": str(col3_value),
                    "col4": row_data.get('col4', ''),
                    "col5": row_data.get('col5', ''),
                    "col6": col6_value,
                    "col7": row_data.get('col7', ''),
                    "suffix": row_data.get('suffix'),
                })
            return result_rows

        output = {
            "type": func_type,
            "func_name": block.get('func_name', fb_key),
        }

        if func_type == 'simple':
            if not block.get('rows'):
                return None
            output["rows"] = prepare_rows(block['rows'])

        elif func_type == 'complex':
            subs = [
                {"subtitle": sub_func.get('subtitle', ''),
                 "rows": prepare_rows(sub_func['rows'])}
                for sub_func in block.get('sub_functions', [])
                if sub_func.get('rows')
            ]
            # Собственные уставки блока (источник fsu-information)
            if block.get('rows'):
                output["rows"] = prepare_rows(block['rows'])
            if not subs and not output.get("rows"):
                return None
            output["sub_functions"] = subs
        else:
            Logger.warning(f"Неизвестный тип блока {func_type} для {fb_key}")
            return None

        return output
