from docx.shared import Cm, Inches
from docx.oxml.shared import OxmlElement, qn
from docx.shared import Pt
import docx
from docx import Document
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from docx.enum.table import WD_ROW_HEIGHT_RULE

import os, sys
import json

from pathlib import Path

from .dropdowns import add_formatted_dropdown2, add_formatted_dropdown3, add_formatted_dropdown2_10pt

def _add_multiline_text_to_cell(cell, text):
    if not text:
        text = ""
    # Преобразуем экранированный \n в настоящий перевод строки
    text = text.replace('\\n', '\n')
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    parts = text.split('\n')
    for i, part in enumerate(parts):
        if i > 0:
            paragraph.add_run().add_break()
        paragraph.add_run(part)

def set_table_borders(table):
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '8')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'auto')
        tblBorders.append(border)
    
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        table._tbl.insert(0, tblPr)
    tblPr.append(tblBorders)

def set_vertical_cell_direction(cell: _Cell, direction: str):
    assert direction in ("tbRl", "btLr")
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    textDirection = OxmlElement('w:textDirection')
    textDirection.set(qn('w:val'), direction)
    tcPr.append(textDirection)

def set_repeat_table_header(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), "true")
    trPr.append(tblHeader)
    return row

def set_cell_vertical_alignment(cell, align="center"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcValign = OxmlElement('w:vAlign')
    tcValign.set(qn('w:val'), align)
    tcPr.append(tcValign)

def set_cell_border(cell: _Cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('start', 'top', 'end', 'bottom', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))

####################################################################################
################################ ТАБЛИЦА ДЛЯ УСТАВОК ХАЛЕЗОВ 07-08-25 ##############################
####################################################################################

table_settings = (Inches(0.25), Inches(1.6), Inches(1.1), Inches(1.6), Inches(0.45), Inches(0.45), Inches(1.4), Inches(1), Inches(1), Inches(1), Inches(1))

def add_table_settings(doc):
    table = doc.add_table(rows=5, cols=11)
    table.style = 'Сетка таблицы51'
    table.allow_autofit = False
    set_table_borders(table)

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '№'
    hdr_cells[1].text = 'Наименование'
    hdr_cells[3].text = 'Значение / Диапазон'
    hdr_cells[4].text = 'Ед. изм.'
    hdr_cells[5].text = 'Шаг'   
    hdr_cells[6].text = 'Значение по умолчанию'
    hdr_cells[7].text = 'Группы уставок'

    for i in range(0,10):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    hdr_cells[1].text = 'ПО ЮС'
    hdr_cells[2].text = 'ИЧМ'
    hdr_cells[7].text = '1'
    hdr_cells[8].text = '2'
    hdr_cells[9].text = '3'
    hdr_cells[10].text = '4'    
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[7].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[8].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[9].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[10].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[2].cells
    tag = f'for row in func.settings'
    hdr_cells[2].text = '{%tr '+ tag + ' %}'
    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = '{{ loop.index }}'
    hdr_cells[1].text = '{{ row[0] }}'
    hdr_cells[2].text = '{{ row[1] }}'    
    hdr_cells[3].text = '{{ row[2]  }}'
    hdr_cells[4].text = '{{ row[3] }}'
    hdr_cells[5].text = '{{ row[4] }}'
    hdr_cells[6].text = '{{ row[5] }}'
    hdr_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[6].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[4].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,9):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 3).merge(table.cell(1, 3))
    table.cell(0, 4).merge(table.cell(1, 4))
    table.cell(0, 5).merge(table.cell(1, 5))
    table.cell(0, 6).merge(table.cell(1, 6))
    table.cell(0, 7).merge(table.cell(0, 10))
    table.cell(2, 0).merge(table.cell(2, 10))
    table.cell(4, 0).merge(table.cell(4, 10))

    for row in table.rows:
        for idx, width in enumerate(table_settings):
            row.cells[idx].width = width

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    return table

####################################################################################
################################ КОНЕЦ ТАБЛИЦА ДЛЯ УСТАВОК #########################
####################################################################################

####################################################################################
############################ ТАБЛИЦА ДЛЯ МАТРИЦЫ ДИСКРЕТНЫХ ВХОДОВ ###############
####################################################################################

table_mtrx_ins = (Inches(2), Inches(4))

def add_table_mtrx_ins(doc, inputs, controls=[]):
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Стиль6'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Дискретный вход'
    hdr_cells[1].text = 'Назначенный сигнал'

    for i in range(0,2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    tag = r'for i in range(1, input_module[1]|int+1)'
    hdr_cells[0].text = '{%tr '+ tag + ' %}'

    hdr_cells = table.rows[2].cells
    hdr_cells[0].text = 'Дискретный вход '+'{{ loop.index }}'

    par1 = hdr_cells[1].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par1,
        inputs_choices=inputs,
        controls_choices=controls,
    )
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(3, 0).merge(table.cell(3, 1))

    for row in table.rows:
        for idx, width in enumerate(table_mtrx_ins):
            row.cells[idx].width = width
    return table 

####################################################################################
############################ КОНЕЦ ТАБЛИЦА ДЛЯ МАТРИЦЫ ДИСКРЕТНЫХ ВХОДОВ ###########
####################################################################################

####################################################################################
############################ ТАБЛИЦА ДЛЯ МАТРИЦЫ ВЫХОДНЫХ РЕЛЕ #####################
####################################################################################

table_mtrx_outs = (Inches(2), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.7))

def add_table_mtrx_outs(doc, statuses, controls=[]):
    table = doc.add_table(rows=5, cols=6)
    table.style = 'Стиль6'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Выходное реле'
    hdr_cells[1].text = 'Назначенные сигналы'

    for i in range(0,6):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    hdr_cells[1].text = '1'
    hdr_cells[2].text = '2'
    hdr_cells[3].text = '3'
    hdr_cells[4].text = '4'
    hdr_cells[5].text = '5'    
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[2].cells
    tag = r'for i in range(1, output_module[1]|int+1)'
    hdr_cells[2].text = '{%tr '+ tag + ' %}'

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = 'Реле '+'{{ loop.index }}'

    par1 = hdr_cells[1].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par1,
        inputs_choices=statuses,
        controls_choices=controls,
    )
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par2 = hdr_cells[2].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par2,
        choices=statuses,
    )
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par3 = hdr_cells[3].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par3,
        choices=statuses,
    )
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par4 = hdr_cells[4].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par4,
        choices=statuses,
    )
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par5 = hdr_cells[5].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par5,
        choices=statuses,
    )
    hdr_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[4].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,6):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(0, 5))
    table.cell(2, 0).merge(table.cell(2, 4))
    table.cell(4, 0).merge(table.cell(4, 4))

    for row in table.rows:
        for idx, width in enumerate(table_mtrx_outs):
            row.cells[idx].width = width
    return table 

####################################################################################
############################ КОНЕЦ ТАБЛИЦА ДЛЯ МАТРИЦЫ ВЫХОДНЫХ РЕЛЕ ###############
####################################################################################

####################################################################################
######################## ТАБЛИЦА ДЛЯ СВЕТОДИОДОВ УСОВЕРШЕНСТВОВАННАЯ ###############
####################################################################################

table_leds_new = (Inches(1.7), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5), Inches(1.5))

def add_table_leds_new(doc, statuses, plates_data):
    table = doc.add_table(rows=5, cols=7)
    table.style = 'Стиль6'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Светодиод'
    hdr_cells[1].text = 'Режим работы'
    hdr_cells[2].text = 'Назначенный сигнал 1'    
    hdr_cells[3].text = 'Назначенный сигнал 2' 
    hdr_cells[4].text = 'Назначенный сигнал 3' 
    hdr_cells[5].text = 'Назначенный сигнал 4' 
    hdr_cells[6].text = 'Назначенный сигнал 5' 

    for i in range(0,7):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    tag = f'for i in range(1, 17)'
    hdr_cells[0].text = '{%tr '+ tag + ' %}'

    hdr_cells = table.rows[2].cells
    hdr_cells[0].text = 'Светодиод '+'{{ loop.index }}' + ' (красный)'
    hdr_cells_row2 = table.rows[3].cells
    hdr_cells_row2[0].text = 'Светодиод '+'{{ loop.index }}'  + ' (зеленый)'

    choices = ["С фиксацией"]
    par2 = hdr_cells[1].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par2,
        choices=choices,
        default='Без фиксации')
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par1 = hdr_cells[2].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par1,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par21 = hdr_cells_row2[1].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par21,
        choices=choices,
        default='Без фиксации')
    hdr_cells_row2[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER 

    par12 = hdr_cells_row2[2].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par12,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells_row2[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par22 = hdr_cells[3].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par22,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER    
    par23 = hdr_cells[4].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par23,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER 
    par24 = hdr_cells[5].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par24,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER 
    par25 = hdr_cells[6].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par25,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells[6].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER 

    par13 = hdr_cells_row2[3].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par13,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков',)
    hdr_cells_row2[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    par14 = hdr_cells_row2[4].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par14,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells_row2[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    par15 = hdr_cells_row2[5].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par15,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells_row2[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    par16 = hdr_cells_row2[6].paragraphs[0]
    add_formatted_dropdown3(
        paragraph=par16,
        inputs_choices=statuses,
        controls_choices = plates_data,
        first_divider= 'Сигналы РЗиА',
        second_divider= 'Сигналы от блоков')
    hdr_cells_row2[6].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[4].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,3):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(1, 0).merge(table.cell(1, 2))
    table.cell(4, 0).merge(table.cell(4, 2))

    for row in table.rows:
        for idx, width in enumerate(table_leds_new):
            row.cells[idx].width = width
    return table 

####################################################################################
################## КОНЕЦ ТАБЛИЦА ДЛЯ СВЕТОДИОДОВ УСОВЕРШЕНСТВОВАННАЯ ###############
####################################################################################

####################################################################################
############################ ТАБЛИЦА ДЛЯ ФУНКЦИОНАЛЬНЫХ КЛАВИШ ###############
####################################################################################

table_fks = (Inches(2), Inches(4))

def add_table_fks(doc, choices_start):
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Стиль6'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Функциональная клавиша'
    hdr_cells[1].text = 'Назначенный сигнал'

    for i in range(0,2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    tag = f'for i in range(1, 17)'
    hdr_cells[0].text = '{%tr '+ tag + ' %}'

    hdr_cells = table.rows[2].cells
    hdr_cells[0].text = 'Функциональная клавиша '+'{{ loop.index }}'

    par1 = hdr_cells[1].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par1,
        choices=choices_start,
    )
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(1, 0).merge(table.cell(1, 1))
    table.cell(3, 0).merge(table.cell(3, 1))

    for row in table.rows:
        for idx, width in enumerate(table_fks):
            row.cells[idx].width = width
    return table

####################################################################################
############################ КОНЕЦ ТАБЛИЦА ДЛЯ ФУНКЦИОНАЛЬНЫХ КЛАВИШ ###############
####################################################################################

#########################################  НОВАЯ  ###################################
################################ ТАБЛИЦА ДЛЯ ДИСКРЕТНЫХ ВХОДОВ ВЫХОДОВ  #############
#####################################################################################

table_binaries = (Inches(0.28), Inches(1.23), Inches(1.4), Inches(1.5), Inches(0.55), Inches(0.45), Inches(0.9), Inches(1.05))

def add_table_binaries(doc, tag = 'for row in items'):
    table = doc.add_table(rows=4, cols=8)
    table.style = 'Сетка таблицы51'
    table.allow_autofit = False
    set_table_borders(table)

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '№'
    hdr_cells[1].text = 'Описание'
    hdr_cells[2].text = 'Наименование'
    hdr_cells[3].text = 'Значение / Диапазон'
    hdr_cells[4].text = 'Ед. изм.'
    hdr_cells[5].text = 'Шаг'
    hdr_cells[6].text = 'Значение по умолчанию'
    hdr_cells[7].text = 'Уставка'
    for i in range(0,8):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    hdr_cells[2].text = '{%tr '+ tag + ' %}'
    hdr_cells = table.rows[2].cells
    hdr_cells[0].text = '{{ loop.index }}'
    hdr_cells[1].text = '{{ row[0] }}'
    hdr_cells[2].text = '{{ row[1] }}'
    hdr_cells[3].text = '{{ row[2]  }}'
    hdr_cells[4].text = '{{ row[3] }}'
    hdr_cells[5].text = '{{ row[4] }}'
    hdr_cells[6].text = '{{ row[5] }}'
    hdr_cells[7].text = ''

    hdr_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[6].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[7].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,8):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(1, 0).merge(table.cell(1, 7))
    table.cell(3, 0).merge(table.cell(3, 7))

    for row in table.rows:
        for idx, width in enumerate(table_binaries):
            row.cells[idx].width = width

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    return table

####################################################################################
######## КОНЕЦ ТАБЛИЦА ДЛЯ ДИСКРЕТНЫХ ФХОДОВ ВЫХОДОВ НОВАЯ #########################
####################################################################################

####################################################################################
################################ ТАБЛИЦА ДЛЯ РЕГИСТРАЦИИ ###########################
####################################################################################

table_reg = (Inches(4.5), Inches(1.5), Inches(1.6), Inches(1.6), Inches(1.6))

def add_table_reg(doc, tag = 'for row in fsu.get_statuses()'):
    table = doc.add_table(rows=5, cols=5)
    table.style = 'Стиль6'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Параметр'
    hdr_cells[2].text = 'Журнал событий регистрация'
    hdr_cells[3].text = 'Осциллограф пуск'
    hdr_cells[4].text = 'Осциллограф регистрация'
    for i in range(0,5):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])

    hdr_cells = table.rows[1].cells
    hdr_cells[0].text = 'Наименование'
    hdr_cells[1].text = 'Обозначение ФСУ'
    hdr_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[2].cells
    hdr_cells[2].text = '{%tr '+ tag + ' %}'

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = '{{ row[0] }}'
    hdr_cells[1].text = '{{ row[1] }}'

    choices_start = ["Не выполняется", "По переднему фронту", "По заднему фронту", "По любому изменению"]
    par3 = hdr_cells[2].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par3,
        choices=choices_start,
    )
    hdr_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    par2 = hdr_cells[3].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par2,
        choices=choices_start,
    )
    hdr_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    choices_reg = ["Выведено", "Введено"]
    par1 = hdr_cells[4].paragraphs[0]
    add_formatted_dropdown2(
        paragraph=par1,
        choices=choices_reg,
    )
    hdr_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells = table.rows[4].cells
    hdr_cells[0].text = '{%tr endfor %}'

    set_repeat_table_header(table.rows[1])
    for i in range(0,5):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'

    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 2).merge(table.cell(1, 2))
    table.cell(0, 3).merge(table.cell(1, 3))
    table.cell(0, 4).merge(table.cell(1, 4))
    table.cell(2, 0).merge(table.cell(2, 4))
    table.cell(4, 0).merge(table.cell(4, 4))

    for row in table.rows:
        for idx, width in enumerate(table_reg):
            row.cells[idx].width = width

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    return table    

####################################################################################
################################ КОНЕЦ ТАБЛИЦА ДЛЯ РЕГИСТРАЦИИ #####################
####################################################################################

####################################################################################
############################ ФИНАЛЬНАЯ ТАБЛИЦА С ПОДПИСЯМИ СОСТАВИТЕЛЯ ###############
####################################################################################

table_final = (Inches(3), Inches(3))

def add_table_final(doc):
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Стиль5'
    table.allow_autofit = False

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'ФИО составителя:'

    hdr_cells = table.rows[1].cells
    hdr_cells[0].text = 'Номер и дата составления:'

    hdr_cells = table.rows[2].cells
    hdr_cells[0].text = 'Дата выдачи:'

    hdr_cells = table.rows[3].cells
    hdr_cells[0].text = 'Дата окончания:'

    table.allow_autofit = False
    table.autofit = False
    table.style = 'Стиль5'

    for row in table.rows:
        for idx, width in enumerate(table_final):
            row.cells[idx].width = width
            row.height = Pt(20)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    return table 

table_settings_core4 = (
    Inches(0.25),
    Inches(2.0),
    Inches(1.8),
    Inches(0.5),
    Inches(0.5),
    Inches(1.5),
    Inches(1),
    Inches(1),
    Inches(1),
    Inches(1)
) 

def add_table_settings_core4(doc):
    table = doc.add_table(rows=2, cols=10)
    table.style = 'Сетка таблицы51'
    table.allow_autofit = False
    
    tbl_pr = table._tbl.xpath('./w:tblPr')[0]
    for elem in tbl_pr.xpath('./w:tblLayout'):
        tbl_pr.remove(elem)
    tbl_pr.append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '№'
    hdr_cells[1].text = 'Наименование'
    hdr_cells[2].text = 'Значение / Диапазон'
    hdr_cells[3].text = 'Ед. изм.'
    hdr_cells[4].text = 'Шаг'   
    hdr_cells[5].text = 'Значение по умолчанию'
    hdr_cells[6].text = 'Группы уставок'

    for i in range(0, 7):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        try:
            set_cell_vertical_alignment(hdr_cells[i], align="center")
        except:
            pass
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    set_repeat_table_header(table.rows[0])

    hdr_cells_2 = table.rows[1].cells
    hdr_cells_2[6].text = '1'
    hdr_cells_2[7].text = '2'
    hdr_cells_2[8].text = '3'
    hdr_cells_2[9].text = '4'
    
    for idx in [6, 7, 8, 9]:
        hdr_cells_2[idx].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(1, 2))
    table.cell(0, 3).merge(table.cell(1, 3))
    table.cell(0, 4).merge(table.cell(1, 4))
    table.cell(0, 5).merge(table.cell(1, 5))
    table.cell(0, 6).merge(table.cell(0, 9))

    for col_idx, width in enumerate(table_settings_core4):
        try:
            table.columns[col_idx].width = width
            table.cell(0, col_idx).width = width
        except IndexError:
            pass

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    return table

####################################################################################
############################ ТАБЛИЦА ДЛЯ МАТРИЦЫ ДИСКРЕТНЫХ ВХОДОВ ###############
####################################################################################

TABLE_WIDTHS_MTRX_INS_CORE4 = (Inches(2), Inches(4))

def add_table_mtrx_ins_core4(doc, slot_name, inputs_list, sigs, di_sigs):
    if not inputs_list:
        return None

    num_rows = len(inputs_list) + 1
    table = doc.add_table(rows=num_rows, cols=2)
    table.style = 'Стиль6'
    table.allow_autofit = False

    try:
        tbl_pr = table._tbl.xpath('./w:tblPr')[0]
        tbl_layout = parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
        tbl_pr.append(tbl_layout)
    except Exception:
        pass

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Дискретный вход'
    hdr_cells[1].text = 'Назначенный сигнал'

    for i in range(2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        if 'set_cell_vertical_alignment' in globals():
            set_cell_vertical_alignment(hdr_cells[i], align="center")

    if 'set_repeat_table_header' in globals():
        set_repeat_table_header(table.rows[0])

    for idx, input_desc in enumerate(inputs_list):
        row_idx = idx + 1
        row_cells = table.rows[row_idx].cells
        
        row_cells[0].text = input_desc
        p_left = row_cells[0].paragraphs[0]
        p_left.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        par_right = row_cells[1].paragraphs[0]
        if 'add_formatted_dropdown3' in globals():
            add_formatted_dropdown3(
                paragraph=par_right,
                inputs_choices=sigs,
                controls_choices=di_sigs,
            )
        else:
            par_right.text = "[Нет сигнала]"
        par_right.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for row in table.rows:
        try:
            row.cells[0].width = TABLE_WIDTHS_MTRX_INS_CORE4[0]
            row.cells[1].width = TABLE_WIDTHS_MTRX_INS_CORE4[1]
        except Exception:
            pass

    return table

TABLE_WIDTHS_OUTS = (Inches(2), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.7), Inches(1.7))

def add_table_mtrx_outs_core4(doc, outputs_list, sigs_list):
    if not outputs_list:
        return
    
    num_rows = len(outputs_list) + 2
    table = doc.add_table(rows=num_rows, cols=6)
    table.style = 'Стиль6'
    table.allow_autofit = False

    try:
        tbl_pr = table._tbl.xpath('./w:tblPr')[0]
        tbl_layout = parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
        tbl_pr.append(tbl_layout)
    except Exception:
        pass

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Выходное реле'
    hdr_cells[1].text = 'Назначенные сигналы'

    for i in range(0, 6):
        p = hdr_cells[i].paragraphs[0]
        try: p.style = 'ДОК Таблица Заголовок'
        except: pass
        if 'set_cell_vertical_alignment' in globals():
            set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    hdr_cells2 = table.rows[1].cells
    hdr_cells2[1].text = '1'
    hdr_cells2[2].text = '2'
    hdr_cells2[3].text = '3'
    hdr_cells2[4].text = '4'
    hdr_cells2[5].text = '5'
    
    for i in range(1, 6):
        p = hdr_cells2[i].paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        try: p.style = 'ДОК Таблица Заголовок'
        except: pass
        if 'set_cell_vertical_alignment' in globals():
            set_cell_vertical_alignment(hdr_cells2[i], align="center")

    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(0, 5))

    for idx, output_name in enumerate(outputs_list):
        row_idx = idx + 2
        row_cells = table.rows[row_idx].cells
        
        row_cells[0].text = output_name
        row_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        for col_idx in range(1, 6):
            par = row_cells[col_idx].paragraphs[0]
            par.clear()
            if 'add_formatted_dropdown2' in globals():
                add_formatted_dropdown2(
                    paragraph=par,
                    choices=sigs_list,
                )
            else:
                par.text = ""
            par.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    if 'set_repeat_table_header' in globals():
        set_repeat_table_header(table.rows[0])
        set_repeat_table_header(table.rows[1])

    for row in table.rows:
        for idx, width in enumerate(TABLE_WIDTHS_OUTS):
            if idx < len(row.cells):
                row.cells[idx].width = width

    return table

####################################################################################
######################## ТАБЛИЦА ДЛЯ СВЕТОДИОДОВ УСОВЕРШЕНСТВОВАННАЯ ###############
####################################################################################

TABLE_WIDTHS_LEDS = (
    Inches(1.2),
    Inches(1.2),
    Inches(1.0),
    Inches(1.5),
    Inches(1.5),
    Inches(1.5),
    Inches(1.5),
    Inches(1.5)
)

def add_table_leds_new_core4(doc, statuses, led_count=16):
    total_rows = 1 + led_count
    table = doc.add_table(rows=total_rows, cols=8)
    table.style = 'Стиль6'
    table.allow_autofit = False

    tbl_pr = table._tbl.xpath('./w:tblPr')
    if tbl_pr:
        tbl_pr[0].append(
            parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
        )

    hdr_cells = table.rows[0].cells
    headers = [
        'Светодиод', 
        'Режим работы', 
        'Цвет', 
        'Назначенный сигнал 1',    
        'Назначенный сигнал 2', 
        'Назначенный сигнал 3', 
        'Назначенный сигнал 4', 
        'Назначенный сигнал 5'
    ]
    
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        try:
            set_cell_vertical_alignment(hdr_cells[i], align="center")
        except NameError:
            pass 

    set_repeat_table_header(table.rows[0]) 

    for row_idx in range(1, total_rows):
        row = table.rows[row_idx]
        cells = row.cells
        
        cells[0].text = f'Светодиод {row_idx}'
        p_name = cells[0].paragraphs[0]
        p_name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        par_mode_data = cells[1].paragraphs[0]
        add_formatted_dropdown2(
            paragraph=par_mode_data,
            choices=["С фиксацией"],
            default='Без фиксации'
        )
        par_mode_data.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        par_color_data = cells[2].paragraphs[0]
        add_formatted_dropdown2(
            paragraph=par_color_data,
            choices=['Зеленый'],
            default='Красный'
        )
        par_color_data.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        for col_idx in range(3, 8):
            par_sig_data = cells[col_idx].paragraphs[0]
            add_formatted_dropdown2(
                paragraph=par_sig_data,
                choices=statuses
            )
            par_sig_data.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for row in table.rows:
        for idx, width in enumerate(TABLE_WIDTHS_LEDS):
            if idx < len(row.cells):
                row.cells[idx].width = width

    return table

####################################################################################
############################ ТАБЛИЦА ДЛЯ ФУНКЦИОНАЛЬНЫХ КЛАВИШ ###############
####################################################################################

TABLE_WIDTHS_KFS = (Inches(2), Inches(4))

def add_table_fks_core4(doc, choices, key_count=16):
    total_rows = 1 + key_count
    table = doc.add_table(rows=total_rows, cols=2)
    table.style = 'Стиль6'
    table.allow_autofit = False

    tbl_pr = table._tbl.xpath('./w:tblPr')
    if tbl_pr:
        tbl_pr[0].append(
            parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
        )

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Функциональная клавиша'
    hdr_cells[1].text = 'Назначенный сигнал'

    for i in range(2):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        try:
            set_cell_vertical_alignment(hdr_cells[i], align="center")
        except NameError:
            pass

    set_repeat_table_header(table.rows[0])

    for row_idx in range(1, total_rows):
        row = table.rows[row_idx]
        cells = row.cells
        
        cells[0].text = f'Функциональная клавиша {row_idx}'
        p_name = cells[0].paragraphs[0]
        p_name.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        par_sig = cells[1].paragraphs[0]
        add_formatted_dropdown2(
            paragraph=par_sig,
            choices=choices
        )
        par_sig.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for row in table.rows:
        for idx, width in enumerate(TABLE_WIDTHS_KFS):
            if idx < len(row.cells):
                row.cells[idx].width = width

    return table

#########################################  НОВАЯ  ###################################
################################ ТАБЛИЦА ДЛЯ ДИСКРЕТНЫХ ВХОДОВ ВЫХОДОВ  #############
#####################################################################################

table_binaries4 = (Inches(0.28), Inches(1.23), Inches(1.4), Inches(1.5), Inches(0.55), Inches(0.45), Inches(0.9), Inches(1.05))

def add_table_binaries_core4(doc, data_rows):
    if not data_rows:
        return None
    
    table = doc.add_table(rows=1 + len(data_rows), cols=8)
    table.style = 'Сетка таблицы51'
    table.allow_autofit = False
    set_table_borders(table)

    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    
    hdr_cells = table.rows[0].cells
    headers = ['№', 'Описание', 'Обозначение ФСУ', 'Значение / Диапазон', 
               'Ед. изм.', 'Шаг', 'Значение по умолчанию', 'Уставка']
    
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    set_repeat_table_header(table.rows[0])

    for row_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[row_idx + 1].cells
        
        row_cells[0].text = str(row_idx + 1)
        row_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Используем функцию для поддержки \n в описании
        _add_multiline_text_to_cell(row_cells[1], str(row_data[0]) if row_data[0] else '')
        row_cells[2].text = str(row_data[1]) if row_data[1] else ''
        row_cells[3].text = str(row_data[2]) if row_data[2] else ''
        row_cells[4].text = str(row_data[3]) if row_data[3] else ''
        row_cells[5].text = str(row_data[4]) if row_data[4] else ''
        row_cells[6].text = str(row_data[5]) if row_data[5] else ''
        row_cells[7].text = ''
        
        row_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        row_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        row_cells[5].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        row_cells[6].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        row_cells[7].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for row in table.rows:
        for idx, width in enumerate(table_binaries4):
            if idx < len(row.cells):
                row.cells[idx].width = width

    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    return table

####################################################################################
################################ ТАБЛИЦА ДЛЯ РЕГИСТРАЦИИ ###########################
####################################################################################

table_reg4 = (Inches(4.0), Inches(2.0), Inches(1.6), Inches(1.6), Inches(1.6))

def add_table_reg_core4(doc, data_rows):
    if not data_rows:
        return None
    
    table = doc.add_table(rows=2 + len(data_rows), cols=5)
    table.style = 'Стиль7'
    table.allow_autofit = False
    
    table._tbl.xpath('./w:tblPr')[0].append(
        parse_xml(r'<w:tblLayout xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:type="fixed"/>')
    )
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Параметр'
    hdr_cells[2].text = 'Журнал событий регистрация'
    hdr_cells[3].text = 'Осциллограф пуск'
    hdr_cells[4].text = 'Осциллограф регистрация'
    for i in range(5):
        p = hdr_cells[i].paragraphs[0]
        p.style = 'ДОК Таблица Заголовок'
        set_cell_vertical_alignment(hdr_cells[i], align="center")
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    
    hdr_cells = table.rows[1].cells
    hdr_cells[0].text = 'Наименование'
    hdr_cells[1].text = 'Обозначение ФСУ'
    hdr_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    hdr_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 2).merge(table.cell(1, 2))
    table.cell(0, 3).merge(table.cell(1, 3))
    table.cell(0, 4).merge(table.cell(1, 4))
    
    choices_reg = ["Введено"]
    choices_osc = ["По переднему фронту", "По заднему фронту", "По любому изменению"]
    
    for row_idx, row_data in enumerate(data_rows):
        row_cells = table.rows[row_idx + 2].cells
        
        col1, col2, type = row_data
        
        row_cells[0].text = str(col1) if col1 else ''
        row_cells[0].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        
        row_cells[1].text = str(col2) if col2 else ''
        row_cells[1].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        
        add_formatted_dropdown2_10pt(
            paragraph=row_cells[2].paragraphs[0],
            choices=choices_osc,
            default="Не выполняется"
        )
        row_cells[2].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        if type != 3:
            row_cells[3].text = 'Не выполняется'
        else:
            add_formatted_dropdown2_10pt(
                paragraph=row_cells[3].paragraphs[0],
                choices=choices_osc,
                default="Не выполняется"
            )        
        row_cells[3].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        if type != 3:
            row_cells[4].text = 'Выведено'
        else:        
            add_formatted_dropdown2_10pt(
                paragraph=row_cells[4].paragraphs[0],
                choices=choices_reg,
                default="Выведено"
            )
        row_cells[4].paragraphs[0].alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    for row in table.rows:
        for idx, width in enumerate(table_reg4):
            if idx < len(row.cells):
                row.cells[idx].width = width
    
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    
    set_repeat_table_header(table.rows[0])
    set_repeat_table_header(table.rows[1])

    return table